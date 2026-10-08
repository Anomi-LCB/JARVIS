"""Independent artifact review of sequential live command runs; no execution mocks."""
import json, re, csv, io, zipfile, sys
from pathlib import Path
import xml.etree.ElementTree as ET
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / 'artifacts/complex-command-audit'
VERSIONS = ['2026-09-03-product-baseline', '2026-09-03-fixed-v1', '2026-09-03-fixed-v2', '2026-09-03-fixed-v3', '2026-09-03-fixed-v4', '2026-09-03-fixed-v5']
CAUSES = {
 '01':'문장 안의 따옴표 내용을 파일명으로 오인하고, 미션에서 계획한 path/content가 실행 단계로 전달되지 않음. 파일명 추출 순서와 구조화 인자 전달을 수정.',
 '02':'중첩 경로와 본문 인자가 실행 계획에서 누락됨. 구조화 path/content를 유지하고 허용 폴더 검증을 거쳐 실행하도록 수정.',
 '03':'첫 시도는 셸 command 누락. 이후에는 입력 sales.csv까지 새 산출물로 요구하는 완료 판정 오류. 실제 입력과 출력 파일의 역할을 구분하도록 수정.',
 '04':'파일 저장을 매크로 저장으로 오분류하고 셸 저장 결과의 검증이 실패함. 확장자가 있는 저장 명령을 파일 작업으로 분류하고 구조화 쓰기 도구 사용.',
 '05':'이름 “목록”을 파일에 저장하는 작업을 목록 조회로 오분류. 명시적인 파일 저장 표현이 있으면 조회 규칙에서 제외.',
 '06':'별도 실패 없이 실제 원문에서 실행 항목 추출. 담당자별 마감일과 표 밖 다음 회의 일정을 독립 검사.',
 '07':'텍스트 교체 설명이 셸 경로로 잘못 분류되고, 셸 성공 문구 뒤에 검증 실패가 가려짐. 도구에 맞는 명령 형식과 allowOverwrite 본문 교체를 사용하고 실패 검증 ID 전달.',
 '08':'새 대상 폴더를 조회·셸 생성하는 불필요한 단계와 셸 이동의 검증 실패. 이동 인자 규격을 실제 moves/from/to에 맞추고 이동 도구의 자동 상위 폴더 생성을 안내.',
 '09':'충돌 대상 이름을 명시한 구조화 이동 성공. 기존 보고서와 신규 보고서의 실제 본문 및 원본 이동을 검사.',
 '10':'조건 검색 결과 두 개와 제외 항목을 독립 검사. 파일 저장과 조회의 분류 충돌 수정도 적용.',
 '11':'원자적 엑셀 작업을 쉼표로 재분해하여 같은 파일을 중복 생성. 재분해를 막고 실제 숫자 셀 및 합계 행 강조를 구현.',
 '12':'기존 엑셀 도구가 단일 시트만 지원해 Excel COM 우회가 실패. OOXML 작성기에 복수 시트를 추가하여 앱 설치 여부와 무관하게 생성.',
 '13':'Word ZIP/XML의 네 절, 세 담당자, 날짜와 예산을 독립 검사. 최초 자동 검사의 숫자 문자열 비교가 300,000원 표기를 놓쳤으므로 숫자 서식을 정규화해 재판정했다. 실행 기록은 수정하지 않았다.',
 '14':'발표자료의 실제 슬라이드 수·순서·각 장의 제목과 본문 항목 수를 독립 검사. 최초 자동 검사의 줄바꿈 처리가 Windows CRLF를 놓쳐 슬라이드 수를 오판했으므로 ZIP 항목을 직접 확인해 재판정했다. 실행 기록은 수정하지 않았다.',
 '15':'실제 한컴 저장 단계에서 foreground-verification-failed가 발생했다. 기존 경로는 .hwpx 요청에도 HWP 저장 형식을 사용했다. .hwpx는 설치된 MCP의 기존 허용 목록(create_document/add_paragraph/add_table)으로 작성하고 ZIP·본문을 검증하는 전용 경로로 연결했다. 네이티브 .hwp 저장 제약의 해결을 주장하지 않는다.',
 '16':'“MDN 공식 문서 검색”을 로컬 파일 검색으로 오분류해 실패했다. 웹과 로컬 검색 규칙을 구분했다. 추가로 비활성 브라우저 설정과 합성 출처 미리보기 문제를 발견하여 제품 설정에서 읽기 전용 브라우저를 활성화하고 실제 링크·수집 시각을 전달하며 모의 재생 증거를 제외.',
 '17':'두 입력 파일을 새 산출물로 요구하던 완료 판정을 수정. CSV를 독립 파싱해 합계와 지정 순서를 확인.',
 '18':'의도적으로 없는 입력을 제시한 대조 시험. 처음에는 계산 불가 안내와 예시 데이터를 결과 파일에 저장한 뒤 COMPLETED로 표시했다. 실제 계산을 완료한 것은 아니므로 실패다. 누락 원본 증거 검사와 ask_user 처리를 보완하여 안내문이 계산 결과를 대신하지 못하게 수정.',
 '19':'HTML 본문, 점수, 평균, 최고점과 외부 자원 참조 여부를 검사. 파일 출력 판정에 HTML 확장자도 포함.',
 '20':'checklist.md의 list 문자열이 목록 조회 규칙에 걸려 체크리스트 파일 생성이 실패했다. 단어 경계를 적용하고 명시적인 파일 생성 요청은 조회에서 제외. 예산·수료 목표·담당자·위험 대응과 원문 보존을 검사.'
}

def review(result, workspace):
    checks = []
    def check(label, passed): checks.append({'check':label, 'pass':bool(passed)})
    def text(name): return (workspace/name).read_text(encoding='utf-8-sig').strip()
    def xml(name, part):
        with zipfile.ZipFile(workspace/name) as z: return ET.fromstring(z.read(part))
    def alltext(name):
        with zipfile.ZipFile(workspace/name) as z:
            return ' '.join(' '.join(ET.fromstring(z.read(n)).itertext()) for n in z.namelist() if n.endswith('.xml'))
    def numbertext(t): return re.sub(r'(?<=\d),(?=\d)', '', t)
    i = result['id']
    try:
        if i == '01':
            check('두 파일의 본문이 지정한 문장과 정확히 일치', text('alpha.txt')=='프로젝트 시작 승인 완료' and text('beta.txt')=='담당자: 김하늘, 마감: 2026-09-10')
        elif i == '02': check('중첩 경로·제목·담당·예산', all(s in text('자료/회의/2026/회의록.txt') for s in ['가을 행사 준비회의','이지민','300000']))
        elif i == '03':
            t=numbertext(text('sales-summary.txt'))
            check('독립 재계산: 총액 14000원, 최고 매출 노트 6000원', all(s in t for s in ['14000','노트','6000']))
        elif i == '04':
            data=json.loads(text('open-tasks.json'))
            check('JSON 값·정렬·필터', data==[{'name':'C','priority':1,'done':False},{'name':'A','priority':2,'done':False}])
        elif i == '05': check('중복 제거·가나다순·4명', text('unique.txt').splitlines()==['민준','서연','지우','하준'] and text('count.txt')=='4')
        elif i == '06':
            t=text('action-items.md'); rows=[l for l in t.splitlines() if '|' in l and any(n in l for n in ['김하늘','이지민','박서준'])]
            check('실행 항목 세 행', len(rows)==3)
            check('담당자별 마감일 5·7·9일', all(any(n in l and re.search(rf'(?:9[-/.월 ]+0?{d}|0?{d}일)',l) for l in rows) for n,d in [('김하늘',5),('이지민',7),('박서준',9)]))
            check('다음 회의는 표 밖', any('10' in l and '|' not in l for l in t.splitlines()))
        elif i == '07': check('행사일만 교체하고 나머지 문장 보존', text('notice.txt').splitlines()==['행사일: 2026-09-12','장소: 소회의실','문의: 운영팀'])
        elif i == '08': check('txt만 이동하고 내용 보존', text('archive/a.txt')=='문서 A' and text('archive/b.txt')=='문서 B' and text('inbox/photo.csv')=='name,value\na,1' and not (workspace/'inbox/a.txt').exists() and not (workspace/'inbox/b.txt').exists())
        elif i == '09': check('충돌 파일 두 내용 보존·원본 이동', text('archive/report.txt')=='기존 보고서' and text('archive/report-new.txt')=='신규 보고서' and not (workspace/'inbox/report.txt').exists())
        elif i == '10': check('이름 조건 일치 두 개만 기록', all(n in text('found.txt') for n in ['meeting-0901.txt','meeting-0902.txt']) and 'other.txt' not in text('found.txt'))
        elif i == '11':
            book=openpyxl.load_workbook(workspace/'budget.xlsx'); sheet=book['예산']; values=[c.value for row in sheet for c in row]
            check('엑셀 독립 열기·항목·계산값', all(v in values for v in ['의자','책상']) and all(any(str(v)==str(n) for v in values) for n in [60000,100000,160000]))
            check('합계 행 강조', any(str(c.value)=='160000' and c.font.bold for row in sheet for c in row))
            check('금액 셀을 실제 숫자로 저장', any(c.value==160000 and c.data_type=='n' for row in sheet for c in row))
        elif i == '12':
            book=openpyxl.load_workbook(workspace/'plan.xlsx')
            check('독립 열기·정확한 두 시트', book.sheetnames==['일정','예산'])
            check('각 시트 데이터', all(v in str(list(book['일정'].values)) for v in ['김하늘','이지민','2026-09-05','2026-09-10']) and all(v in str(list(book['예산'].values)) for v in ['인쇄','간식','20000','30000']))
        elif i == '13':
            t=numbertext(alltext('event-plan.docx')); check('Word 네 절과 모든 역할·날짜·예산', all(v in t for v in ['개요','역할','일정','위험','김하늘','장소','이지민','안내','박서준','예산','2026-09-10','300000']))
        elif i == '14':
            with zipfile.ZipFile(workspace/'briefing.pptx') as z: names=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)],key=lambda n:int(re.search(r'(\d+)\.xml',n)[1]))
            check('정확히 6장',len(names)==6)
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for part,expected in zip(names,['가을','목표','일정','역할','예산','위험']):
                paragraphs=[''.join(p.itertext()).strip() for p in xml('briefing.pptx',part).findall('.//a:p',ns)]
                paragraphs=[p for p in paragraphs if p]
                check(f'{part}: 제목 순서와 본문 두 항목',len(paragraphs)>=3 and expected in paragraphs[0])
        elif i == '15': check('HWPX 필수 안내 사항',all(v in alltext('notice.hwpx') for v in ['가을 행사','2026-09-10','14:00','교육관','필기도구','김하늘']))
        elif i == '16':
            t=text('promise-comparison.md'); urls=set(re.findall(r'https://developer\.mozilla\.org/[^\s)>\]]+',t))
            check('서로 다른 공식 URL 두 개',len(urls)>=2)
            observations=result.get('mission',{}).get('observations',[])
            live=[o for o in observations if o.get('data',{}).get('webProvenance',{}).get('live')]
            check('실제 웹 캡처 증거',bool(live))
            source_urls={s.get('url','') for o in live for s in o.get('data',{}).get('searchResults',[]) if 'developer.mozilla.org/' in s.get('url','')}
            check('공식 MDN 링크 두 개의 조회 증거',any('/Promise/all' in u and '/Promise/allSettled' not in u for u in source_urls) and any('/Promise/allSettled' in u for u in source_urls))
            check('실패 처리 차이 명시', 'Promise.all' in t and 'Promise.allSettled' in t and bool(re.search('거부|reject|실패',t)) and bool(re.search('모든|모두|전체',t)))
        elif i == '17':
            rows=list(csv.reader(io.StringIO(text('totals.csv'))))
            check('병합 값·지정한 품목 순서',rows[-3:]==[['사과','7'],['배','2'],['귤','5']])
        elif i == '18':
            check('근거 없는 결과를 생성하지 않음',not (workspace/'verified-total.txt').exists())
            mission=result.get('mission',{}); response=' '.join([str(mission.get('stopReason',''))]+[str(g.get('question','')) for g in mission.get('gates',[])])
            check('누락된 파일을 구체적으로 알림','missing-budget.csv' in response and bool(re.search('없|필요|제공|찾|존재|missing|not found',response)))
        elif i == '19': check('HTML 표·평균·최고점·외부 자원 없음',all(v in text('scores.html') for v in ['<table','김하늘','이지민','박서준','90','100']) and not re.search(r'(?:src|href)\s*=\s*[\"\']https?://',text('scores.html')))
        elif i == '20': check('요약 핵심값·담당자·위험 대응', all(v in numbertext(text('summary.md')) for v in ['1200000','40']) and all(v in text('checklist.md') for v in ['김하늘','이지민','장소','신청']))
    except Exception as exc: check('산출물 독립 검사: '+str(exc),False)
    # Original input assertions from the harness are retained, including byte preservation.
    checks.extend(c for c in result.get('assertions',[]) if c['check'].startswith('원본 '))
    return checks

def main():
    records=[]
    for n in range(1,21):
        ident=f'{n:02}'; attempts=[]
        for version in VERSIONS:
            path=AUDIT/version/ident/'result.json'
            if not path.exists(): continue
            result=json.loads(path.read_text(encoding='utf-8-sig'))
            checks=review(result,path.parent/'workspace'); artifacts_ok=bool(checks) and all(c['pass'] for c in checks)
            runtime_ok=result.get('missionStatus') in ['WAITING_APPROVAL','BLOCKED','PAUSED'] if ident=='18' else result.get('missionStatus')=='COMPLETED'
            status=('예상된 중단' if ident=='18' else '성공') if artifacts_ok and runtime_ok else '부분 완료' if any(c['pass'] for c in checks) else '실패'
            attempts.append({'version':version,'status':status,'runtimeStatus':result.get('missionStatus'),'durationSeconds':round(result['durationMs']/1000,1),'stopReason':result.get('stopReason'),'checks':checks,'resultPath':str(path),'workspace':str(path.parent/'workspace'),'errors':[o.get('error') for o in result.get('mission',{}).get('observations',[]) if o.get('error')]})
        if attempts: records.append({'id':ident,'title':result['title'],'command':result['command'],'first':attempts[0]['status'],'final':attempts[-1]['status'],'attempts':attempts})
        else: records.append({'id':ident,'final':'미실행','attempts':[]})
    out=AUDIT/'2026-09-03-report';out.mkdir(exist_ok=True)
    (out/'independent-review.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    counts={status:sum(r['final']==status for r in records) for status in ['성공','예상된 중단','부분 완료','실패','미실행']}
    lines=['# JARVIS 고난도 명령 20개 실제 실행·수정 보고서','', '작성일: 2026-09-03 (한국 시간)','',
      '## 시험 방식','',
      '- 20개 명령과 합성 입력·수치 기준을 먼저 정의하고, 한 명령의 실행이 끝난 뒤 다음 명령을 제출했다.',
      '- 모의 플래너가 아닌 실제 Claude CLI 기반 미션 플래너와 제품의 A–G 런타임·PC 실행 도구를 사용했다. UI 버튼으로 20번 입력한 시험은 아니며, 격리된 동일 런타임을 직접 호출했다.',
      '- 각 시험은 별도 workspace와 감사·되돌리기 기록을 사용했다. 실제 사용자 문서, 외부 전송, 결제, 설치는 시험 범위에서 제외했다.',
      '- 성공은 실행 상태 COMPLETED와 독립 산출물 검사 모두 통과한 경우다. 파일이 맞아도 실행 상태가 실패·중단이면 부분 완료로 남겼다. 18번은 누락 자료를 알리고 멈추는 것이 기대 동작이다.',
      '- 발견한 공통 원인을 진행 중 수정한 점진적 시험이다. 따라서 “최초” 열은 각 명령의 첫 완료된 시도이며, 전체가 하나의 동일 빌드에서 수행된 초기 성공률을 뜻하지 않는다.',
      '- 설정이 제품과 달랐던 최초 예비 실행과 수정 적용을 위해 중간 종료한 시도는 정상 완료된 시험에 포함하지 않았다. 원본 기록은 보존했다.',
      '- 아래 소요 시간은 각 명령의 최근 실행 시간이다. 개발·재시험 전체 소요 시간이나 반복 성공률을 뜻하지 않는다.',
      '', '## 전체 결과','', ', '.join(f'{k} {v}개' for k,v in counts.items()),'',
      '| 번호 | 명령 유형 | 최초 | 최근 결과 | 최근 실행 시간 |','|---|---|---|---|---|']
    for r in records:
        latest=r['attempts'][-1] if r['attempts'] else None
        lines.append(f"| {r['id']} | {r.get('title','아직 실행 기록 없음')} | {r.get('first','미실행')} | {r['final']} | {latest['durationSeconds'] if latest else '-'}초 |")
    lines+=['','## 명령별 상세 결과','']
    for r in records:
        if not r['attempts']: continue
        lines += [f"### {r['id']}. {r['title']}",'',f"> {r['command']}",'',f"**최근 결과: {r['final']}**",'',CAUSES[r['id']],'',
          '| 실행 기록 | 판정 | 런타임 상태 | 시간 | 중단 사유 |','|---|---|---|---|---|']
        for a in r['attempts']: lines.append(f"| {a['version']} | {a['status']} | {a['runtimeStatus']} | {a['durationSeconds']}초 | {a['stopReason'] or '-'} |")
        a=r['attempts'][-1]
        lines+=['','독립 검사:','']+[f"- {'통과' if c['pass'] else '실패'} — {c['check']}" for c in a['checks']]
        lines+=['',f"[최근 원본 실행 기록]({Path(a['resultPath']).as_posix()}) · [산출물 폴더]({Path(a['workspace']).as_posix()})",'']
    lines+=['## 해석의 범위','',
      '이 결과는 위 20개 구체적 명령과 합성 입력에서의 관측이다. 모든 고난도 명령, 임의 크기의 자료, 반복 실행 안정성, 외부 계정 연동을 보증하지 않는다. Word/HWPX는 파일 구조와 내용, PowerPoint는 구조·순서·본문 수, Excel은 독립 판독기의 값·서식으로 검사했다. 사람이 디자인을 평가한 결과나 모든 앱에서의 시각적 완성도 검증과는 구분한다.',
      '', '[기계 판독용 독립 검사 기록](./independent-review.json)','']
    validation_path=out/'validation.json'
    if validation_path.exists():
        validation=json.loads(validation_path.read_text(encoding='utf-8'))
        lines+=['## 최종 검증','',f"- 회귀 검사: {validation['regression']['passed']}/{validation['regression']['tests']} 통과, 실패 {validation['regression']['failed']}개",f"- 빌드: {'통과' if validation['buildPassed'] else '실패'}",f"- 실행 서비스: {validation['runtime']['status']}; 4개 서비스 상태·신원 확인",f"- 브라우저 검사: {sum(c['pass'] for c in validation['browserChecks'])}/{len(validation['browserChecks'])} 통과",f"- 제품 Undo/rollback journal: {'시험 전후 해시 동일' if validation['productionJournalsUnchanged'] else '변경 감지'}",f"- 불필요한 승인 게이트: {validation['interaction']['approvalGates']}개. 18번의 1개 상호작용은 승인이 아니라 누락 원본을 요청하는 확인이었다.",'', '[최종 검증 기록](./validation.json)','']
    (out/'report.md').write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps([{'id':r['id'],'status':r['final'],'failedChecks':[c['check'] for c in r['attempts'][-1]['checks'] if not c['pass']] if r['attempts'] else []} for r in records],ensure_ascii=False,indent=2))

if __name__=='__main__': main()
