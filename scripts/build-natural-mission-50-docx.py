from __future__ import annotations

import argparse
import importlib.util
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


def build(source: Path, target: Path):
    helper_path = Path(__file__).with_name('build-impossible-workflow-30-docx.py')
    spec = importlib.util.spec_from_file_location('report_layout', helper_path)
    layout = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(layout)
    report = json.loads(source.read_text(encoding='utf-8-sig'))
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.32)
    for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[style_name]
        style.font.name = layout.FONT
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), layout.FONT)
        style.font.size = Pt(11)
    layout.add_footer(section)

    def body(text, size=11):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.line_spacing = 1.15
        layout.set_font(paragraph.add_run(str(text)), size=size)
        return paragraph

    title = doc.add_paragraph(report['title'], style='Title')
    for run in title.runs:
        layout.set_font(run, size=23, bold=True)
    body(f"작성일 {report['date']}   작성 Codex   수신 JARVIS 사용자", 9)
    for text in report['overview']:
        body(text)
    layout.add_heading(doc, '검증 결과')
    layout.add_table(doc, ['항목', '확인된 결과'], report['metrics'], [5.0, 12.75], font_size=10)

    for heading, paragraphs in report.get('sections', []):
        layout.add_heading(doc, heading)
        for text in paragraphs:
            body(text)
    layout.add_heading(doc, '실패 원인과 실제 수정')
    for item in report['repairs']:
        layout.add_heading(doc, item['name'], 2)
        body('발견 내용  ' + item['problem'])
        body('원인  ' + item['cause'])
        body('수정 및 검증  ' + item['solution'])
    layout.add_heading(doc, '추가 검증과 성능')
    for text in report['verification']:
        body(text)
    layout.add_heading(doc, '권장하는 다음 개선')
    for item in report['recommendations']:
        layout.add_heading(doc, f"{item['priority']} {item['name']}", 2)
        body(item['reason'])
        body('구체적인 다음 작업  ' + item['next'])
        if item.get('source'):
            body('근거  ' + item['source'], 9)
    layout.add_heading(doc, '검증 범위와 남은 한계')
    for text in report['limits']:
        body(text)

    doc.add_page_break()
    layout.add_heading(doc, '미션 50개 개별 명령과 결과')
    body('다음 문장은 실제 실행기에 전달한 원문입니다. 입력은 재현 가능한 합성 자료이며 각 사례의 원본 기록과 결과 파일은 보고서 폴더에 보존했습니다. 성공은 제품 실행 완료와 별도 내용 검증을 모두 통과한 경우입니다.')
    last_category = None
    for case in report['cases']:
        if case['category'] != last_category:
            layout.add_heading(doc, case['category'])
            last_category = case['category']
        heading = re.sub(r'[^\w\s]', ' ', f"{case['id']} {case['title']}")
        layout.add_heading(doc, re.sub(r'\s+', ' ', heading), 2)
        metadata = body(f"복합 지시 {case['clauseCount']}개   수정 전 {case['before']}   수정 후 {case['after']}   실행 {case['durationMs'] / 1000:.3f}초", 9.5)
        metadata.paragraph_format.keep_with_next = True
        body(case['command'], 10.5)
        body('판정  ' + case['reason'], 10)
        body('결과  ' + '; '.join(case['outputs']), 9.5)
        if case.get('sample'):
            body('내용 예시  ' + case['sample'], 9)

    layout.add_heading(doc, '증거 위치')
    for item in report['evidence']:
        body(item['label'], 10)
        body(item['path'], 8.5)
    for element in [doc._element, doc.styles._element]:
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(json.dumps({'output': str(target.resolve()), 'cases': len(report['cases'])}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('target', type=Path)
    args = parser.parse_args()
    build(args.source, args.target)
