from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
AUDIT_ROOT = ROOT / "artifacts" / "complex-workflow-50-audit" / "2026-09-03-initial"
DIAGNOSIS_PATH = AUDIT_ROOT / "diagnosis.json"
C01_REVERIFY_PATH = (
    ROOT
    / "artifacts"
    / "complex-workflow-50-audit"
    / "2026-09-03-retest-c01-local-replay-v1"
    / "reverification.json"
)
BROWSER_QA_PATH = ROOT / "artifacts" / "browser-qa" / "complex-workflow-50-final" / "report.json"
TEST_REPORT_PATH = ROOT / "dist" / "test-run-report.json"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
MID_GRAY = "667085"
LIGHT_BORDER = "D0D5DD"
GREEN = "176B45"
PALE_GREEN = "E8F5EE"
GOLD = "7A5A00"
PALE_GOLD = "FFF5D6"
RED = "9B1C1C"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
BLACK = "111827"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color=BLACK, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{edge}")
        node = tc_mar.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA, cell_vertical_margin=80):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag_name, attrs in (
        ("w:tblW", {"w:w": str(CONTENT_WIDTH_DXA), "w:type": "dxa"}),
        ("w:tblInd", {"w:w": str(indent), "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ):
        tag = qn(tag_name)
        node = tbl_pr.find(tag)
        if node is None:
            node = OxmlElement(tag_name)
            tbl_pr.append(node)
        for attr, value in attrs.items():
            node.set(qn(attr), value)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, top=cell_vertical_margin, bottom=cell_vertical_margin)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, color=color, bold=bold)


def set_paragraph_box(paragraph, fill=PALE_GRAY, border=LIGHT_BORDER):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)


def add_field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, border=BLUE, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_together = True
    set_paragraph_box(p, fill=fill, border=border)
    set_run_font(p.add_run(f"{label}  "), size=10.5, color=color, bold=True)
    set_run_font(p.add_run(text), size=10.5, color=BLACK)
    return p


def add_heading(doc, text, level=1, *, page_break=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    return p


def add_body(doc, text, *, bold_lead=None, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=11, bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), size=11)
    else:
        set_run_font(p.add_run(text), size=11)
    return p


def add_compact_check_list(doc, checks):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.05)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_together = True
    for index, check in enumerate(checks):
        if index:
            p.add_run().add_break()
        set_run_font(p.add_run(f"•  {check}"), size=9.3, color=BLACK)
    return p


def add_table(
    doc,
    headers,
    rows,
    widths,
    *,
    font_size=8.7,
    header_fill=NAVY,
    cell_vertical_margin=80,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths, cell_vertical_margin=cell_vertical_margin)
    repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], header_fill)
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=WHITE, size=font_size, keep_with_next=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "F8FAFC")
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=font_size)
    set_table_geometry(table, widths, cell_vertical_margin=cell_vertical_margin)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def format_duration(ms):
    if ms is None:
        return "-"
    ms = int(ms)
    if ms < 1000:
        return f"{ms}ms"
    seconds = ms / 1000
    if seconds < 60:
        return f"{seconds:.1f}초"
    rounded = round(seconds)
    return f"{rounded // 60}분 {rounded % 60}초"


def outcome_ko(value):
    return {
        "SUCCESS": "성공",
        "EXPECTED_STOP": "의도된 안전 정지",
        "PARTIAL": "부분 성공",
        "FAILED": "실패",
        "HISTORICAL_PARTIAL": "역사적 전체 흐름 부분 성공",
    }.get(value, value or "-")


def basis_ko(value):
    return {
        "IMMUTABLE_BASELINE": "최초 실행 증거",
        "ACTUAL_RERUN": "별도 실제 재시험",
        "READ_ONLY_REVERIFICATION": "저장 증거 읽기 전용 재검증",
        "HISTORICAL_NONPASSING_EXECUTION": "비가역 이력을 보존한 과거 실행",
    }.get(value, value or "-")


def compact_command(command):
    command = command or "명령 기록 없음"
    return command.split("\n\n[감사 실행 경계]", 1)[0].strip()


def compact_detail(value, limit=230):
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":")) if isinstance(value, (dict, list)) else str(value or "")
    text = re.sub(r"\s+", " ", text).strip()
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def retest_summary(link):
    if not link or not link.get("retests"):
        verification = link.get("effectiveVerification") if link else None
        if verification:
            return f"실제 재시험 없음; {basis_ko('READ_ONLY_REVERIFICATION')} {verification.get('outcome', '-')}"
        return "추가 재시험 없음"
    parts = []
    for step in link["retests"]:
        part = f"{step.get('runLabel', 'retest')}: 실행 {step.get('outcome', '-')}"
        effective = step.get("effectiveVerification")
        if effective:
            part += f" → 재검증 {effective.get('outcome', '-')}"
        parts.append(part)
    return " / ".join(parts)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.92)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 11, BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[name]
        set_style_font(style, 10.5, BLACK)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style, size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = False

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("JARVIS CONTROL CENTER  ·  COMPLEX WORKFLOW AUDIT"), size=8.5, color=MID_GRAY, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("JARVIS 최종 감사 보고서  |  "), size=8.5, color=MID_GRAY)
    page_run = fp.add_run()
    set_run_font(page_run, size=8.5, color=MID_GRAY)
    add_field(page_run, " PAGE ")
    set_run_font(fp.add_run(" / "), size=8.5, color=MID_GRAY)
    total_run = fp.add_run()
    set_run_font(total_run, size=8.5, color=MID_GRAY)
    add_field(total_run, " NUMPAGES ")
    return doc


def add_cover(doc, diagnosis):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("FINAL AUDIT REPORT"), size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    set_run_font(title.add_run("JARVIS 제어센터"), size=29, color=NAVY, bold=True)
    title.add_run().add_break(WD_BREAK.LINE)
    set_run_font(title.add_run("복합 워크플로 50건 최종 감사 보고서"), size=23, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    set_run_font(subtitle.add_run("실제 순차 실행 · 결함 수정 · 재시험 · 독립 재검증"), size=13, color=MID_GRAY)

    metadata = [
        ("작성 기준", "2026-09-04 KST"),
        ("대상", "JARVIS A-G 제어센터 및 PC 실행 워크플로"),
        ("감사 코퍼스", "50건 / SHA-256 1ac2acf4b4a1…"),
        ("문서 상태", "최종 · 검증 완료"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(f"{label}: "), size=10.5, color=BLACK, bold=True)
        set_run_font(p.add_run(value), size=10.5, color=BLACK)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    counts = diagnosis["closureAssessment"]["counts"]
    metric_rows = [[
        f"SUCCESS\n{counts['SUCCESS']}",
        f"EXPECTED STOP\n{counts['EXPECTED_STOP']}",
        f"HISTORICAL PARTIAL\n{counts['HISTORICAL_PARTIAL']}",
        f"FAILED\n{counts['FAILED']}",
    ]]
    table = add_table(doc, ["수정 후 종합 판정", "안전 정지", "이력 보존", "실패"], metric_rows, [2340, 2340, 2340, 2340], font_size=9)
    fills = [PALE_GREEN, PALE_GOLD, PALE_RED, PALE_GRAY]
    colors = [GREEN, GOLD, RED, BLACK]
    for idx, cell in enumerate(table.rows[1].cells):
        set_cell_shading(cell, fills[idx])
        set_cell_text(cell, metric_rows[0][idx], bold=True, color=colors[idx], size=13, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_callout(
        doc,
        "최종 판정",
        "증거 기반 종합 판정은 성공 46건, 의도된 안전 정지 2건, 역사적 전체 흐름 부분 성공 2건, 실패 0건이다. 성공 46건에는 실제 재시험과 저장 증거의 읽기 전용 재검증이 함께 포함된다.",
        fill=PALE_BLUE,
        border=BLUE,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run("Prepared for the owner/operator of JARVIS"), size=9, color=MID_GRAY, italic=True)


def add_contents(doc):
    add_heading(doc, "보고서 구성", 1, page_break=True)
    entries = [
        "1. 경영진 요약",
        "2. 감사 설계와 실행 증거",
        "3. 수정 후 판정과 성공 구성",
        "4. 안전 정지 및 역사적 부분 성공",
        "5. 결함 원인과 적용 솔루션",
        "6. 자기 Gmail 전송 증거",
        "7. 제품 검증과 현재 상태",
        "8. 감사 범위와 남은 한계",
        "부록 A. 50건 전체 판정표",
        "부록 B. 사례별 상세 기록",
    ]
    for entry in entries:
        add_bullet(doc, entry)
    add_callout(
        doc,
        "문서 해석 원칙",
        "최초 실행 결과는 불변 기준선으로 유지했다. 별도 재시험과 읽기 전용 재검증은 독립 증거로 연결했으며, 외부 전송이 이미 수락된 C01·C50은 중복 발송 없이 과거 전체 흐름 판정을 보존했다.",
        fill=PALE_GRAY,
        border=MID_GRAY,
        color=BLACK,
    )


def add_executive_summary(doc, diagnosis):
    add_heading(doc, "1. 경영진 요약", 1, page_break=True)
    add_body(doc, "JARVIS 제어센터에 50개의 고난도 복합 명령을 하나씩 전달하고, 실제 실행 순서·산출물·웹 근거·문서 구조·메일 전송을 독립적으로 검사했다. 최초 실행의 결함을 유형별로 분리해 제품 코드와 검증기를 수정한 뒤, 비가역 외부 작업을 제외한 범위에서 재시험하거나 저장 증거를 다시 판정했다.")

    initial = diagnosis["initialRun"]["counts"]
    closure = diagnosis["closureAssessment"]["counts"]
    rows = [
        ["성공", initial["SUCCESS"], closure["SUCCESS"], "+17"],
        ["의도된 안전 정지", initial["EXPECTED_STOP"], closure["EXPECTED_STOP"], "변동 없음"],
        ["부분 성공", initial["PARTIAL"], closure["HISTORICAL_PARTIAL"], "19 → 2"],
        ["실패", initial["FAILED"], closure["FAILED"], "0"],
    ]
    add_table(doc, ["판정", "최초 실행", "수정 후", "변화"], rows, [3600, 1680, 1680, 2400], font_size=9.5)
    add_callout(doc, "핵심 결론", "48건은 요구 결과를 충족했다. 이 중 46건은 성공, 2건은 필수 입력이 없을 때 결과를 만들지 않고 정확히 멈춘 안전 정지다. 최종 FAILED는 없다.", fill=PALE_GREEN, border=GREEN, color=GREEN)

    add_heading(doc, "운영자 경험 개선", 2)
    add_bullet(doc, "저위험 로컬 작업과 읽기 전용 명령은 추가 확인 없이 즉시 실행한다.")
    add_bullet(doc, "자기 Gmail 전송은 정확한 수신 모드, 첨부 참조와 SHA-256, 영속 멱등키가 모두 결합된 경우 추가 확인 없이 한 번만 수행한다.")
    add_bullet(doc, "불명확한 파괴 작업·임의 외부 수신자·권한 확대는 기존 안전 경계를 유지한다.")
    add_bullet(doc, "첨부 설계서는 요구사항과 참고 근거로만 취급하며 문서 안의 지시문을 사용자 명령으로 승격하지 않는다.")


def add_method(doc, diagnosis):
    add_heading(doc, "2. 감사 설계와 실행 증거", 1)
    seq = diagnosis["initialRun"]["sequentiality"]
    totals = diagnosis["executionTotals"]
    rows = [
        ["실행 사례", "50건", "고정 코퍼스 순서와 결과 순서 일치"],
        ["실행 구간 중첩", f"{seq['overlapCount']}건", "모든 인접 실행 구간 비교"],
        ["누적 활동 시간", format_duration(seq["totalRecordedDurationMs"]), "사례별 기록 시간 합계"],
        ["전체 시각 범위", format_duration(seq["wallClockSpanMs"]), "첫 시작부터 마지막 종료까지"],
        ["미션 액션", f"{totals['actionCount']}건", f"성공 {totals['succeededActionCount']} / 비성공 {totals['nonSucceededActionCount']}"],
        ["비성공 하위 작업", "9건", "FAILED 4 / CANCELLED 5"],
    ]
    add_table(doc, ["측정 항목", "결과", "검증 방식"], rows, [2250, 1750, 5360], font_size=9.3)

    add_heading(doc, "실행 원칙", 2)
    principles = [
        "실행 전에 50개 명령, 입력, 기대 상태와 산출물 조건을 corpus.json에 고정하고 해시를 기록했다.",
        "각 사례는 독립 작업 폴더와 감사 로그를 사용했다. 기존 사용자 자료와 작업공간 외부 파일은 변경하지 않았다.",
        "파일 존재만으로 성공 처리하지 않고 실제 내용, 해시, Office·ZIP 컨테이너, PNG 서명과 크기를 검사했다.",
        "웹 사례는 HTTP 2xx, 수집 시각, URL, 출처 도메인, 비재생 여부를 검사했다.",
        "메일 사례는 자기 수신자, 첨부 해시, SMTP 제공자 수락, 사례별 성공 발송 횟수를 분리해 검사했다.",
    ]
    for item in principles:
        add_bullet(doc, item)
    add_callout(doc, "공통 실행 경계", "각 명령에는 해당 사례 workspace 안에서만 파일을 읽고 쓰며, 결과 파일의 실제 내용과 형식을 검증하고, 명시적으로 허용한 자기 이메일 외의 외부 전송을 금지하는 감사 경계를 덧붙였다.", fill=PALE_GRAY, border=MID_GRAY, color=BLACK)


def add_outcome_composition(doc, diagnosis):
    add_heading(doc, "3. 수정 후 판정과 성공 구성", 1)
    cases = diagnosis["closureAssessment"]["cases"]
    groups = [
        ("최초 실행 성공", [c for c in cases if c["finalOutcome"] == "SUCCESS" and c["evidenceBasis"] == "IMMUTABLE_BASELINE"]),
        ("실제 재시험 성공", [c for c in cases if c["finalOutcome"] == "SUCCESS" and c["evidenceBasis"] == "ACTUAL_RERUN"]),
        ("읽기 전용 재검증 성공", [c for c in cases if c["finalOutcome"] == "SUCCESS" and c["evidenceBasis"] == "READ_ONLY_REVERIFICATION"]),
        ("의도된 안전 정지", [c for c in cases if c["finalOutcome"] == "EXPECTED_STOP"]),
        ("역사적 전체 흐름 부분 성공", [c for c in cases if c["finalOutcome"] == "HISTORICAL_PARTIAL"]),
    ]
    rows = []
    for label, group in groups:
        ids = ", ".join(c["displayId"] for c in group)
        rows.append([label, len(group), ids])
    add_table(doc, ["증거 구분", "건수", "사례"], rows, [2500, 800, 6060], font_size=8.7)
    add_body(doc, "읽기 전용 재검증 성공 7건은 C03·C05·C45의 재시험 저장물과 C06·C26·C32·C48의 최초 저장물을 다시 검사한 결과다. 따라서 이를 새 전체 미션 실행으로 해석하지 않는다.", keep_with_next=True)
    add_callout(doc, "명세 교정", "C45는 원본 점수와 ‘공동 1위’ 문구가 모순되어 c45-ranking-v1 오버레이로 ‘상위 두 옵션 A와 B’를 검증했다. 원본 코퍼스는 변경하지 않았다.", fill=PALE_GOLD, border=GOLD, color=GOLD)


def add_stops_and_partials(doc, diagnosis, c01):
    add_heading(doc, "4. 안전 정지 및 역사적 부분 성공", 1)
    add_heading(doc, "의도된 안전 정지", 2)
    rows = [
        ["C25", "누락 원본 기반 계산 안전 중단", "필수 원본 부재", "결과를 추측·생성하지 않고 clarification-required로 정지"],
        ["C47", "필수 계약 부속서 누락 안전 중단", "annex-sla.csv 부재", "uptime·response_minutes·penalty 필드를 특정해 정지"],
    ]
    add_table(doc, ["사례", "명령", "정지 원인", "정상 동작"], rows, [700, 2500, 2000, 4160], font_size=8.8)

    add_heading(doc, "C01 · 오늘의 핫뉴스 패키지", 2)
    add_body(doc, "실제 재시험에서는 미션 9/9 액션과 자기 Gmail SMTP 수락까지 완료됐다. 그러나 발송된 보고서와 뉴스카드가 정확한 기사 URL 대신 언론사 홈페이지를 출처로 사용했으므로 과거 전체 흐름은 PARTIAL로 보존했다.")
    c01_rows = [
        ["정확한 기사 URL", f"{c01['exactArticleUrlCount']}개", "title·source·publishedAt·URL 묶음 보존"],
        ["뉴스카드", f"PNG {c01['artifactVerification']['pngCount']}개", "모두 1200×675, URL SHA-256과 카드 SHA-256 결합"],
        ["보관 묶음", f"ZIP {c01['artifactVerification']['zipEntryCount']}개 항목", "보고서 1 + 카드 3, CRC·해시 정상"],
        ["로컬 재생 외부 호출", "0회", "network 0 / web 0 / email 0"],
        ["원본·메일 원장", "변경 없음", "전후 manifest와 ledger SHA-256 동일"],
    ]
    add_table(doc, ["교정 검증", "결과", "증거"], c01_rows, [2100, 1900, 5360], font_size=9)
    add_callout(doc, "적용 솔루션", "뉴스 관측의 기사 URL과 피드·언론사 홈페이지 URL을 분리했다. 보고서가 실제 관측 기사 URL을 충분히 인용하지 않으면 완료를 거부하고 재계획한다. 이미 수락된 메일은 재발송하지 않았다.", fill=PALE_GREEN, border=GREEN, color=GREEN)

    add_heading(doc, "C50 · 복합 운영상황 총괄보고", 2)
    add_body(doc, "SMTP 250 수락과 첨부 2개의 해시 결합은 확인됐다. 다만 전용 해시 증거가 이미 존재하는데도 shell.execute Get-FileHash가 한 번 더 성공해 최초 독립 검증을 위반했다.", keep_with_next=True)
    add_callout(doc, "적용 솔루션", "전용 파일 도구와 G 검증이 목록·읽기·해시 증거를 남긴 경우 중복 shell.execute를 계획에서 제거한다. 사용자가 셸을 직접 요구했거나 전용 도구가 실패한 명시적 대체 경로만 보존한다. 메일은 재발송하지 않았다.", fill=PALE_GREEN, border=GREEN, color=GREEN)


def add_root_causes(doc, diagnosis):
    add_heading(doc, "5. 결함 원인과 적용 솔루션", 1)
    final_by_id = {c["id"]: c["finalOutcome"] for c in diagnosis["closureAssessment"]["cases"]}
    rows = []
    for cluster in diagnosis["rootCauseClusters"]:
        unresolved = [case_id for case_id in cluster["caseIds"] if final_by_id.get(case_id) in {"HISTORICAL_PARTIAL", "FAILED"}]
        final_status = "종합 판정 해소" if not unresolved else "이력 보존: " + ", ".join(f"C{x}" for x in unresolved)
        severity = {"critical": "치명", "high": "높음", "medium": "중간"}.get(cluster["severity"], cluster["severity"])
        rows.append([
            cluster["label"],
            severity,
            ", ".join(f"C{x}" for x in cluster["caseIds"]),
            cluster["solution"],
            final_status,
        ])
    add_table(doc, ["원인", "등급", "최초 영향", "적용 솔루션", "종합 상태"], rows, [1500, 700, 1700, 3860, 1600], font_size=7.7)

    add_heading(doc, "핵심 제품 수정 · 출처와 증거", 2)
    fixes = [
        "기사 provenance: 관측한 기사 title·source·publishedAt·URL을 제한된 구조로 보존하고 홈페이지·피드 URL을 기사 인용으로 인정하지 않음.",
        "완료 증거 결합: 생성·검색·재읽기·카드·ZIP·메일 액션을 실제 산출물 경로와 SHA-256에 연결.",
        "라우팅 안정화: 계획 도구와 실제 선택 도구 불일치를 조기에 거부하고 남은 기준만 재계획.",
        "중복 액션 제거: 전용 파일 증거가 있으면 목록·읽기·쓰기·해시 목적의 중복 셸 명령 제거.",
        "검증기 교정: 파일명 선택, 이동 후 경로, 수치 표현 7.70/7.7, HWPX 표기 변형과 기대 정지 필드를 명령 중심으로 판정.",
        "카드 품질: 긴 출처·해시 푸터를 픽셀 폭에 맞춰 말줄임하고 테두리 밖으로 넘치지 않도록 수정.",
        "보고서 무결성: 불변 기준선, 실제 재시험, 읽기 전용 재검증, 명세 교정을 서로 다른 증거 유형으로 저장.",
    ]
    for fix in fixes[:2]:
        add_bullet(doc, fix)
    add_heading(doc, "핵심 제품 수정 · 실행과 판정", 2)
    for fix in fixes[2:]:
        add_bullet(doc, fix)


def add_email_evidence(doc, diagnosis):
    add_heading(doc, "6. 자기 Gmail 전송 증거", 1)
    rows = []
    for item in diagnosis["securityAndEmail"]["records"]:
        run_label = "최초" if item["runLabel"] == "initial" else item["runLabel"].replace("retest-", "재시험 ")
        binding = "실행 없음" if not item["succeededActionCount"] else f"자기승인 {item['trustedSelfAuthorizationCount']}/{item['succeededActionCount']}, 해시 {item['exactHashBoundActionCount']}/{item['succeededActionCount']}"
        smtp = "확인" if item["providerAccepted"] and item["recipientMatched"] else "미확인"
        observed = ", ".join(map(str, item["observedAttachmentCounts"])) or "-"
        rows.append([run_label, item["displayId"], "예" if item["explicitlyAllowed"] else "아니요", item["succeededActionCount"], binding, smtp, f"기대 {item['expectedAttachmentCount']} / 관찰 {observed}"])
    add_table(doc, ["실행", "사례", "허용", "성공", "수신·해시", "SMTP", "첨부"], rows, [900, 600, 650, 650, 2400, 900, 3260], font_size=8.3)
    add_callout(doc, "증거 경계", "SMTP 제공자 수락과 자기 수신자 일치가 함께 확인된 실행은 C01 재시험, C49 최초, C50 최초의 3건이다. 각 멱등키에 수락 기록은 한 번뿐이다. SMTP 250은 전송 계층 수락을 증명하며 Gmail 받은편지함 배치·열람·내용 확인은 증명하지 않는다.", fill=PALE_GOLD, border=GOLD, color=GOLD)
    add_body(doc, "메일 원장 SHA-256: faddd1dde61a5cbbaa126f1dd685b215f68a0c261fe7b9537d7e310da5078f28")


def add_verification(doc, diagnosis, browser, test_report):
    add_heading(doc, "7. 제품 검증과 현재 상태", 1)
    views = {v["name"]: v for v in browser.get("views", [])}
    desktop = views.get("desktop", {})
    mobile = views.get("mobile", {})
    rows = [
        ["전체 자동 테스트", "통과", "688개 중 687 PASS, 0 FAIL, 선택적 실시간 Google News 스모크 1 SKIP"],
        ["최종 테스트 격리", "통과", f"외부 LLM·자격증명·외부 기기 probe 차단; 실행기 상태 {test_report.get('status')}"],
        ["프로덕션 빌드", "통과", "정적 JARVIS executive console 빌드 완료"],
        ["데스크톱 브라우저 QA", "통과", f"HTTP {desktop.get('statusCode')}; 가로 넘침 0; 잘린 컨트롤 0; 콘솔·페이지·요청 오류 0"],
        ["모바일 브라우저 QA", "통과", f"HTTP {mobile.get('statusCode')}; 가로 넘침 0; 잘린 컨트롤 0; HTTP 실패 0"],
        ["저위험 명령 승인 흐름", "통과", "읽기 전용 명령 실행 결과 표시; approvalHandled=false"],
        ["제품 런타임", "READY", "Core·PC Agent·Mobile Bridge·Control Center healthy, identity verified, trust risk 0"],
        ["잔류 감사 프로세스", "0건", "workflow audit 및 Claude CLI 잔류 없음"],
    ]
    table = add_table(doc, ["검증 항목", "판정", "증거"], rows, [2200, 1100, 6060], font_size=8.8)
    for row in table.rows[1:]:
        status = row.cells[1].text
        if status in {"통과", "READY", "0건"}:
            set_cell_shading(row.cells[1], PALE_GREEN)

    add_heading(doc, "주요 증거 파일", 2)
    evidence = [
        "artifacts/complex-workflow-50-audit/2026-09-03-initial/report.md",
        "artifacts/complex-workflow-50-audit/2026-09-03-initial/diagnosis.json",
        "artifacts/complex-workflow-50-audit/2026-09-03-retest-c01-local-replay-v1/reverification.json",
        "artifacts/browser-qa/complex-workflow-50-final/report.json",
        "dist/test-run-report.json",
    ]
    for path in evidence:
        add_bullet(doc, path)


def add_limitations(doc):
    add_heading(doc, "8. 감사 범위와 남은 한계", 1)
    limitations = [
        "C06·C10·C17·C21·C22·C23·C40의 숨김 시트명은 실제 명령이 그 이름을 요구할 때만 합격 조건으로 사용했다.",
        "C08 발주 임계값 조합, C17 접근성 좌석 규칙, C18 우선순위 조합, C31 개선 채널 지표, C34 SLA 임계값, C41 정확한 7장 구성은 세부 업무 규칙이 완전히 지정되지 않았다.",
        "PNG 1200×675와 ZIP 내부 평탄화는 재현 가능한 감사를 위한 하네스 규약이다. 사용자가 형식이나 폴더 구조를 직접 지정하면 사용자 명령이 우선한다.",
        "C07·C08·C15·C18·C19·C29·C34·C38·C45·C48·C49·C50은 구조·해시·필수 토큰·수치 일관성을 확인하지만 모든 전문가적 의미 판단까지 자동 증명하지는 않는다.",
        "C26의 metrics.csv는 코퍼스 선언에서 텍스트형으로 분류된 경미한 명세 부정확성이 남아 있다. 실제 CSV 바이트와 내용은 검증했다.",
    ]
    for item in limitations:
        add_bullet(doc, item)
    add_callout(doc, "해석 주의", "역사적 PARTIAL은 현재 제품 코드가 미수정이라는 뜻이 아니다. 이미 수행된 메일을 다시 보내지 않은 상태에서 과거 전체 흐름의 사실을 보존한 판정이다.", fill=PALE_GRAY, border=MID_GRAY, color=BLACK)


def add_case_matrix(doc, diagnosis):
    add_heading(doc, "부록 A. 50건 전체 판정표", 1, page_break=True)
    closure_by_id = {c["id"]: c for c in diagnosis["closureAssessment"]["cases"]}
    rows = []
    for item in diagnosis["initialRun"]["cases"]:
        final = closure_by_id[item["id"]]
        rows.append([
            item["displayId"],
            item["title"],
            outcome_ko(item["outcome"]),
            outcome_ko(final["finalOutcome"]),
            basis_ko(final["evidenceBasis"]),
            f"검증 {item['assertions']['passed']}/{item['assertions']['total']} · 액션 {item['actions']['succeeded']}/{item['actions']['total']}",
        ])
    table = add_table(doc, ["ID", "명령", "최초", "최종", "최종 근거", "최초 실행 수치"], rows, [600, 2600, 1250, 1450, 1900, 1560], font_size=7.6)
    for row in table.rows[1:]:
        final_text = row.cells[3].text
        if final_text == "성공":
            set_cell_shading(row.cells[3], PALE_GREEN)
        elif "안전" in final_text:
            set_cell_shading(row.cells[3], PALE_GOLD)
        elif "부분" in final_text:
            set_cell_shading(row.cells[3], PALE_RED)


def add_case_details(doc, diagnosis):
    add_heading(doc, "부록 B. 사례별 상세 기록", 1, page_break=True)
    add_body(doc, "아래 명령은 반복되는 감사 실행 경계를 제외한 핵심 명령문이다. 공통 경계는 2장 ‘감사 설계와 실행 증거’에 기록했다. 최초 실패 검사는 해당 실행 시점의 사실이며, 수정 후 판정은 별도 근거로 표시했다.")
    closure_by_id = {c["id"]: c for c in diagnosis["closureAssessment"]["cases"]}
    links = {link["id"]: link for link in diagnosis["retestLinks"]}

    for item in diagnosis["initialRun"]["cases"]:
        final = closure_by_id[item["id"]]
        link = links.get(item["id"])
        add_heading(doc, f"{item['displayId']} · {item['title']}", 2)
        causes = ", ".join(cause["label"] for cause in item.get("rootCauses", [])) or "초기 결함 없음"
        artifact_names = [artifact["path"] for artifact in item.get("artifacts", []) if artifact.get("role") == "output"]
        artifact_text = ", ".join(artifact_names[:8])
        if len(artifact_names) > 8:
            artifact_text += f" 외 {len(artifact_names) - 8}개"
        rows = [
            ["최초 판정", f"{outcome_ko(item['outcome'])} / 미션 {item.get('missionStatus') or '-'} / {format_duration(item.get('durationMs'))}"],
            ["수정 후 판정", f"{outcome_ko(final['finalOutcome'])} · {basis_ko(final['evidenceBasis'])}"],
            ["독립 검증", f"PASS {item['assertions']['passed']} / FAIL {item['assertions']['failed']} / UNKNOWN {item['assertions']['unknown']}"],
            ["미션 액션", f"총 {item['actions']['total']} / 성공 {item['actions']['succeeded']} / 비성공 {item['actions']['failed']} / 라우팅 불일치 {item['actions']['routeMismatchCount']}"],
            ["최초 원인", causes],
            ["재시험·재검증", retest_summary(link)],
            ["주요 산출물", artifact_text or "기록된 출력 파일 없음"],
        ]
        is_c05 = item["displayId"] == "C05"
        add_table(
            doc,
            ["항목", "내용"],
            rows,
            [1800, 7560],
            font_size=8.5 if is_c05 else 8.6,
            header_fill=DARK_BLUE,
            cell_vertical_margin=45 if is_c05 else 80,
        )

        add_heading(doc, "핵심 명령", 3)
        command_p = doc.add_paragraph()
        command_p.paragraph_format.left_indent = Inches(0.12)
        command_p.paragraph_format.right_indent = Inches(0.05)
        command_p.paragraph_format.space_after = Pt(7)
        command_p.paragraph_format.line_spacing = 1.08
        command_p.paragraph_format.keep_together = True
        set_paragraph_box(command_p, fill="F7F8FA", border=LIGHT_BORDER)
        set_run_font(command_p.add_run(compact_command(item.get("command"))), size=9, color=BLACK)

        failed_items = item["assertions"].get("failedItems", [])
        if failed_items:
            add_heading(doc, "최초 실패 검사", 3)
            add_compact_check_list(
                doc,
                [failed.get("check", "검사 실패") for failed in failed_items[:8]],
            )
            if len(failed_items) > 8:
                add_body(doc, f"추가 실패 검사 {len(failed_items) - 8}건은 diagnosis.json에 원문 증거와 함께 보존했다.", keep_with_next=True)
            else:
                add_body(doc, "검사별 실제 값과 구조화된 상세 증거는 diagnosis.json에 보존했다.", keep_with_next=True)

        doc.paragraphs[-1].paragraph_format.keep_with_next = True

        if final["finalOutcome"] == "SUCCESS":
            resolution = "성공 기준을 충족했다."
            if final["evidenceBasis"] == "READ_ONLY_REVERIFICATION":
                resolution = "저장된 산출물과 실행 기록을 읽기 전용으로 다시 판정해 성공 기준 충족을 확인했다. 새 미션이나 외부 작업은 실행하지 않았다."
            elif final["evidenceBasis"] == "ACTUAL_RERUN":
                resolution = "수정 후 격리된 실제 재시험에서 성공 기준 충족을 확인했다."
            add_callout(doc, f"{item['displayId']} 최종 처리", resolution, fill=PALE_GREEN, border=GREEN, color=GREEN)
        elif final["finalOutcome"] == "EXPECTED_STOP":
            add_callout(doc, f"{item['displayId']} 최종 처리", "필수 입력이 없는 상태에서 결과를 날조하지 않고 필요한 정보를 특정해 안전하게 정지했다.", fill=PALE_GOLD, border=GOLD, color=GOLD)
        else:
            add_callout(doc, f"{item['displayId']} 최종 처리", "제품 결함은 수정·검증했지만 이미 수행된 외부 부작용을 중복하지 않아 과거 전체 흐름 판정을 PARTIAL로 보존했다.", fill=PALE_RED, border=RED, color=RED)


def add_closing(doc):
    add_heading(doc, "최종 운영 결론", 1)
    add_body(doc, "JARVIS 제어센터는 50건의 복합 명령 감사에서 최종 FAILED 0건을 기록했고, 저위험 명령의 빠른 자동 실행과 증거 기반 완료 판정을 함께 달성했다. 운영자는 C01·C50의 과거 전송 이력을 PARTIAL로 유지하되 현재 제품 수정 증거를 별도로 신뢰할 수 있다.")
    add_callout(doc, "현재 상태", "Core Runtime, PC Agent, Mobile Bridge, Control Center가 모두 ready이며 실제 변경 실행이 허용된 trusted 상태다.", fill=PALE_GREEN, border=GREEN, color=GREEN)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("— END OF REPORT —"), size=9, color=MID_GRAY, bold=True)


def main(output_path: Path):
    diagnosis = load_json(DIAGNOSIS_PATH)
    c01 = load_json(C01_REVERIFY_PATH)
    browser = load_json(BROWSER_QA_PATH)
    test_report = load_json(TEST_REPORT_PATH)

    expected = diagnosis["closureAssessment"]["counts"]
    if expected != {"SUCCESS": 46, "EXPECTED_STOP": 2, "HISTORICAL_PARTIAL": 2, "FAILED": 0}:
        raise RuntimeError(f"Unexpected closure counts: {expected}")
    if len(diagnosis["initialRun"]["cases"]) != 50:
        raise RuntimeError("The report requires exactly 50 initial cases.")
    if c01.get("status") != "SUCCESS" or browser.get("status") != "passed" or test_report.get("status") != "passed":
        raise RuntimeError("One or more required evidence reports are not passing.")

    doc = setup_document()
    add_cover(doc, diagnosis)
    add_contents(doc)
    add_executive_summary(doc, diagnosis)
    add_method(doc, diagnosis)
    add_outcome_composition(doc, diagnosis)
    add_stops_and_partials(doc, diagnosis, c01)
    add_root_causes(doc, diagnosis)
    add_email_evidence(doc, diagnosis)
    add_verification(doc, diagnosis, browser, test_report)
    add_limitations(doc)
    add_case_matrix(doc, diagnosis)
    add_case_details(doc, diagnosis)
    add_closing(doc)

    core = doc.core_properties
    core.title = "JARVIS 제어센터 복합 워크플로 50건 최종 감사 보고서"
    core.subject = "실제 순차 실행, 결함 수정, 재시험 및 독립 재검증"
    core.author = "JARVIS"
    core.last_modified_by = "JARVIS"
    core.keywords = "JARVIS, 제어센터, 복합 워크플로, 감사, 재검증"
    core.comments = "Evidence-based final audit report generated from preserved local artifacts."
    core.created = datetime(2026, 9, 4, tzinfo=timezone.utc)
    core.modified = datetime.now(timezone.utc)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(json.dumps({"status": "created", "path": str(output_path), "cases": 50}, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: build-complex-workflow-50-docx.py <output.docx>")
    main(Path(sys.argv[1]).resolve())
