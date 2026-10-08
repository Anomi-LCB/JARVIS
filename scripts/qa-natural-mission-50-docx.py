import json
import re
import sys
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document

root = Path(sys.argv[1]).resolve()
source = root / 'JARVIS_복합명령50_개선결과보고서.docx'
render = root / 'render'
document = Document(source)
text = '\n'.join(p.text for p in document.paragraphs)
cases = re.findall(r'^N\d{2}\s', text, re.M)
assert len(cases) == 50 and len(set(cases)) == 50
assert document.sections[0].page_width.inches == 8.5
assert document.sections[0].page_height.inches == 11
assert document.paragraphs[0].style.name == 'Title'
pdf = pdfium.PdfDocument(str(render / 'report.pdf'))
pages = []
all_text = ''
for index in range(len(pdf)):
    page = pdf[index]
    text_page = page.get_textpage()
    content = text_page.get_text_range()
    all_text += content + '\n'
    image = page.render(scale=100 / 72).to_pil()
    image.save(render / f'page-{index + 1:02}.png')
    boxes = [text_page.get_charbox(i) for i in range(text_page.count_chars())]
    footer_digits = [(text_page.get_text_range(i, 1), box) for i, box in enumerate(boxes)
                     if box[1] < 40 and text_page.get_text_range(i, 1).isdigit()]
    assert ''.join(char for char, box in footer_digits) == str(index + 1)
    assert max(box[1] for char, box in footer_digits) - min(box[1] for char, box in footer_digits) < 0.5
    meaningful = [box for box in boxes if box[2] > box[0] and box[3] > box[1]]
    width, height = page.get_size()
    outside = [box for box in meaningful if box[0] < 0 or box[1] < 0 or box[2] > width + .1 or box[3] > height + .1]
    assert not outside, (index + 1, outside)
    pages.append({'page': index + 1, 'characters': len(content), 'lines': len(content.splitlines()),
                  'textExcerpt': content[:110], 'endExcerpt': content[-120:], 'outOfPageCharacters': len(outside),
                  'pageNumberOnSingleLine': True, 'pageNumberGlyphPositions': footer_digits})
assert all(f'N{index:02}' in all_text for index in range(1, 51))
assert '\ufffd' not in all_text
report = {'pages': len(pdf), 'docxCases': len(cases), 'allCaseIdsPresentInPdf': True,
          'renderer': 'Microsoft Word ExportAsFixedFormat plus bundled pypdfium2; packaged renderer failed because LibreOffice soffice.exe is unavailable',
          'structureStatus': 'PASSED', 'pageSummary': pages}
(render / 'structure-qa.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'pages': len(pdf), 'docxCases': len(cases), 'structureStatus': 'PASSED'}))
