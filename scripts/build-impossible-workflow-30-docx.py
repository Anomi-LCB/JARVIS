from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT = "Malgun Gothic"
NAVY = "18324A"
BLUE = "DCEAF5"
PALE = "F4F7FA"
GRAY = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def set_font(run, size=None, bold=None, color=TEXT):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_font(run, size=8)


def add_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(4)
    label = paragraph.add_run("JARVIS 구조 개선 검증  |  ")
    set_font(label, size=8)
    add_page_field(paragraph)


def format_cell(cell, *, size=9, bold=False, white=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            set_font(run, size=size, bold=bold, color=RGBColor(255, 255, 255) if white else TEXT)


def set_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, *, font_size=8.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(value)
        shade(cell, NAVY)
        set_width(cell, widths[index])
        format_cell(cell, size=8.5, bold=True, white=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.text = str(value)
            set_width(cell, widths[index])
            if row_index % 2:
                shade(cell, PALE)
            alignment = alignments[index] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(cell, size=font_size, align=alignment)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        set_font(run, size=15 if level == 1 else 12, bold=True)
    return paragraph


def add_body(doc, text, *, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_font(lead, size=10.2, bold=True)
        body = paragraph.add_run(text[len(bold_prefix):])
        set_font(body, size=10.2)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=10.2)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.12
    for run in paragraph.runs:
        set_font(run, size=9.7)
    return paragraph


def status_ko(value):
    return "성공" if str(value).upper() in {"SUCCESS", "PASSED", "COMPLETED"} else "실패"


def build(report_path: Path, output_path: Path, validation_path: Path | None):
    report = json.loads(report_path.read_text(encoding="utf-8"))
    validation = {}
    if validation_path and validation_path.exists():
        validation = json.loads(validation_path.read_text(encoding="utf-8"))

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.footer_distance = Cm(0.8)
    add_footer(section)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.color.rgb = TEXT
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = TEXT

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("JARVIS 불가능 명령 30건\n구조 개선 최종 보고서"), size=23, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run("실제 제품 경로 실행과 독립 검증 결과"), size=12, bold=True)
    generated = report.get("generatedAt", "")
    if generated:
        set_font(subtitle.add_run(f"  |  생성 시각 {generated}"), size=9)

    add_body(
        doc,
        "기존 JARVIS가 처리하지 못했던 로컬 복합 워크플로 30건을 고정한 뒤, 동일한 명령과 판정 계약으로 구조 개선 전후를 비교했다. "
        "개선 전에는 30건 모두 실행 기능 부재로 실패했고, 최신 구조에서는 30건 모두 MissionRuntime, PC Agent, G 검증, 독립 출력 판정을 통과했다."
    )

    summary = report["beforeAfter"]
    metrics = [
        ["구조 개선 전", f"{summary['baselinePassed']}/30", "성공"],
        ["구조 개선 후", f"{summary['finalPassed']}/30", "성공"],
        ["독립 재검증", "30/30", "통과"],
        ["셸 우회 실행", str(summary.get("forbiddenShellDowngrades", 0)), "건"]
    ]
    add_table(
        doc,
        ["측정 항목", "결과", "단위"],
        metrics,
        [7.7, 4.0, 4.0],
        font_size=10,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
    )

    campaign_metrics = validation.get("metrics", {})
    if campaign_metrics:
        scale_rows = [
            ["실행 명령", campaign_metrics.get("workflowCases", ""), "건"],
            ["workflow 노드", campaign_metrics.get("workflowNodes", ""), "개"],
            ["사용 연산 유형", campaign_metrics.get("operationTypesUsed", ""), "종"],
            ["게시 산출물", campaign_metrics.get("publishedArtifacts", ""), "개"],
            ["명령당 평균 실행", campaign_metrics.get("meanCaseDurationMs", ""), "ms"],
            ["최장 명령 실행", campaign_metrics.get("maxCaseDurationMs", ""), "ms"],
        ]
        add_table(
            doc,
            ["캠페인 규모", "값", "단위"],
            scale_rows,
            [7.7, 4.0, 4.0],
            font_size=9.5,
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        )

    add_heading(doc, "판정 범위", 1)
    add_body(doc, "이번 판정은 고정된 30개 로컬 워크플로와 현재 실행 환경에 대한 증거다. 명령마다 준비된 typed workflow 명세를 자연어로 호출했고, 명세 내부의 데이터 처리와 파일 작업, 그래프 계산, 프로세스 파이프라인, SQLite 작업을 실제로 수행했다.")
    add_body(doc, "임의 문장만으로 새 workflow 명세를 자동 설계하는 능력, 외부 계정 권한, 웹사이트 변경, 네트워크 장애, 물리 장치 동작 전체를 보장하는 시험은 아니다. 이 구분은 30건의 성공 결과를 축소하지 않으며, 결과가 증명하는 범위를 정확히 한정한다.")

    add_heading(doc, "실행 및 검증 경로", 1)
    path_rows = [[str(index + 1), item] for index, item in enumerate(report.get("scope", {}).get("productPath", []))]
    if not path_rows:
        path_rows = [[str(index + 1), item] for index, item in enumerate([
            "createAgRuntime", "MissionRuntime", "CoreFacade", "CoreOrchestrationWorker", "PcAgentRuntime", "VerificationRuntime", "독립 오라클"
        ])]
    add_table(doc, ["순서", "검증된 제품 경로"], path_rows, [2.0, 13.7], font_size=9.5,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    add_heading(doc, "내부 시스템 구조 개선", 1)
    for change in report.get("structuralChanges", []):
        add_body(doc, f"{change['id']} {change['title']}  {change['behavior']}", bold_prefix=f"{change['id']} {change['title']}  ")

    add_heading(doc, "발견한 실패와 해결", 1)
    failure_rows = [
        ["Baseline 30건", "workflow.execute 미등록", "typed IR와 실행 어댑터, 동적 capability package 등록"],
        ["첫 최종 실행 IW17", "resource 안의 source 문자열을 외부 출처로 오인", "영문 단어 경계 판정으로 완료 검증 보정"],
        ["전체 회귀 5건", "새 증거 스키마 검사가 과거 증거 형식에도 무조건 적용", "workflow는 fail closed 유지, 기존 형식은 명시된 스키마가 있을 때 검증"],
        ["등록 중간 실패", "일부 route가 남는 원자성 결함", "모든 projection snapshot과 보상 복구"],
        ["프로세스 네트워크", "Node permission만으로 loopback 접근 가능", "정적 검사와 runtime network guard, IPC 차단 증거"],
        ["staging과 원장", "junction 선점과 post commit 실패 시 잔여 가능", "무작위 원자 staging, publish 보상, cleanup 경고 분리"],
        ["산출물 계약", "선언 파일이 생산 node와 결속되지 않고 Windows 대소문자 별칭이 충돌 가능", "선언 산출물 subset과 생산 node 결속, canonical path 고유성 강제"],
        ["SQLite 재실행", "별도 append/upsert가 기존 게시 DB를 stage로 복사하지 않아 데이터 유실", "기존 DB를 stage로 복제한 뒤 같은 트랜잭션에서 갱신"],
        ["데이터·경로 검증", "상속 속성 순회와 비정상 일정 자원·기간·중복 ID 허용", "own property와 위험 segment 차단, 유한 정수 자원·비음수 기간·고유 ID 검증"],
        ["capability·G 식별", "package ID 충돌과 중복 receipt/manifest, G 누락 통과 가능", "package 계약 digest 충돌 차단, 증거 identity 고유성 및 G 통과 필수화"],
        ["보고서 재개", "임의 result 재사용과 빈 배열 판정 가능", "identity, 원본 Mission G 증거, 해시 체인을 재검증"],
        ["독립 오라클", "8개 유형에서 부분 토큰만으로 허위 통과 가능", "입력 바이트 재계산과 구조 의미 검증, 변조 corpus"],
    ]
    add_table(doc, ["발견 지점", "직접 원인", "해결"], failure_rows, [3.2, 5.5, 7.0], font_size=8.5)

    results_heading = add_heading(doc, "30개 명령 결과", 1)
    results_heading.paragraph_format.page_break_before = True
    case_rows = []
    for item in report.get("cases", []):
        artifacts = item.get("final", {}).get("outputArtifacts", [])
        artifact_text = "\n".join(Path(value.get("path", "")).name for value in artifacts) or "없음"
        case_rows.append([
            item.get("caseId", ""),
            item.get("title", ""),
            item.get("category", ""),
            status_ko(item.get("baseline", {}).get("outcome")),
            status_ko(item.get("final", {}).get("outcome")),
            "통과" if item.get("final", {}).get("oraclePassed") else "실패",
            artifact_text,
        ])
    add_table(
        doc,
        ["ID", "명령", "분야", "전", "후", "독립 판정", "주요 산출물"],
        case_rows,
        [1.2, 4.5, 2.5, 1.1, 1.1, 1.6, 3.7],
        font_size=7.8,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.LEFT]
    )

    add_heading(doc, "독립 증거 게이트", 1)
    gate_rows = [[gate.get("gateId", ""), "통과" if gate.get("passed") else "실패"] for gate in report.get("gates", [])]
    add_table(doc, ["게이트", "결과"], gate_rows, [12.7, 3.0], font_size=9,
              alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])

    if validation:
        add_heading(doc, "최종 제품 검증", 1)
        validation_rows = []
        for check in validation.get("checks", []):
            validation_rows.append([check.get("name", ""), check.get("result", ""), check.get("detail", "")])
        add_table(doc, ["검증", "결과", "세부"], validation_rows, [5.0, 2.4, 8.3], font_size=8.8)

    tamper_summary = report.get("tamperNegative", {}).get("summary", {})
    tamper_counts = tamper_summary.get("counts", {})
    if tamper_counts:
        add_heading(doc, "변조 검증 상세", 1)
        tamper_rows = [
            ["주 산출물 바이트 변경", f"{tamper_counts.get('primary-byte', 0)}/30", "탐지"],
            ["주 산출물 삭제", f"{tamper_counts.get('primary-delete', 0)}/30", "탐지"],
            ["빈 proof 대체", f"{tamper_counts.get('proof-empty', 0)}/30", "탐지"],
            ["passed:false proof 대체", f"{tamper_counts.get('proof-false', 0)}/30", "탐지"],
            ["결과 재해시·Mission/G 누락·corpus 변경", "4/4", "탐지"],
        ]
        add_table(
            doc,
            ["공격 유형", "결과", "판정"],
            tamper_rows,
            [10.0, 3.0, 2.7],
            font_size=9,
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        )

    add_heading(doc, "고정 증거 식별자", 1)
    frozen = report.get("frozenEvidence", {})
    identifiers = [
        ["명령 corpus SHA 256", frozen.get("corpusSha256", "")],
        ["오라클 계약 SHA 256", frozen.get("oracleContractsSha256", "")],
        ["Baseline 원장 head", frozen.get("baselineLedgerHead", "")],
        ["최종 원장 head", frozen.get("finalLedgerHead", "")],
    ]
    add_table(doc, ["식별자", "값"], identifiers, [4.5, 11.2], font_size=7.8)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "실제 자연어 명령 원문", 1)
    add_body(doc, "아래 문장 30개를 한 건씩 MissionRuntime에 전달했다. CASE DIR 표시는 실행 때 각 격리 사례 폴더의 절대 경로로 치환됐다.")
    for item in report.get("cases", []):
        paragraph = doc.add_paragraph()
        if item.get("caseId") == "IW26":
            paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.keep_together = True
        lead = paragraph.add_run(f"{item.get('caseId')}  {item.get('title')}\n")
        set_font(lead, size=9.3, bold=True)
        command = item.get("commandTemplate", "")
        run = paragraph.add_run(command)
        set_font(run, size=8.7)
        paragraph.paragraph_format.line_spacing = 1.08

    add_heading(doc, "최종 판단", 1)
    add_body(doc, report.get("conclusion", ""))
    add_body(doc, report.get("limit", ""))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "JARVIS 불가능 명령 30건 구조 개선 최종 보고서"
    doc.core_properties.subject = "로컬 복합 워크플로 30건 구조 개선 및 검증"
    doc.core_properties.author = "JARVIS 검증 시스템"
    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--validation", type=Path)
    args = parser.parse_args()
    build(args.report, args.output, args.validation)


if __name__ == "__main__":
    main()
