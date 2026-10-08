from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def main(docx_path: Path, report_path: Path, output_path: Path) -> None:
    issues: list[str] = []
    report = json.loads(report_path.read_text(encoding="utf-8"))
    expected_cases = report.get("cases", [])
    expected_ids = [f"IW{index:02d}" for index in range(1, 31)]

    docx_bytes = docx_path.read_bytes()
    with zipfile.ZipFile(docx_path) as archive:
        corrupt_member = archive.testzip()
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml")
        footer_parts = [archive.read(name) for name in names if re.fullmatch(r"word/footer\d+\.xml", name)]
    if corrupt_member:
        issues.append(f"corrupt ZIP member: {corrupt_member}")
    if any("comment" in name.lower() for name in names):
        issues.append("comment package part remains")

    root = etree.fromstring(document_xml)
    tracked = {tag: len(root.xpath(f".//w:{tag}", namespaces=NS)) for tag in ("ins", "del", "moveFrom", "moveTo")}
    if any(tracked.values()):
        issues.append(f"tracked changes remain: {tracked}")

    doc = Document(docx_path)
    if not (20.9 <= doc.sections[0].page_width.cm <= 21.1 and 29.6 <= doc.sections[0].page_height.cm <= 29.8):
        issues.append("first section is not A4 portrait")

    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    paragraph_text = "\n".join(paragraphs)
    if "JARVIS 불가능 명령 30건\n구조 개선 최종 보고서" not in paragraph_text:
        issues.append("title is missing")
    if "임의의 미래 명령" not in paragraph_text:
        issues.append("scope boundary is missing")

    appendix_ids: list[str] = []
    for paragraph in paragraphs:
        match = re.match(r"^(IW\d{2})\s{2}", paragraph)
        if match:
            appendix_ids.append(match.group(1))
    if appendix_ids != expected_ids:
        issues.append(f"appendix IDs differ: {appendix_ids}")

    if len(expected_cases) != 30:
        issues.append(f"report case count is {len(expected_cases)}, expected 30")
    for case in expected_cases:
        command = case.get("commandTemplate", "")
        if not any(command in paragraph for paragraph in paragraphs):
            issues.append(f"full command is missing: {case.get('caseId')}")

    table_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    expected_gate_ids = [gate.get("gateId") for gate in report.get("gates", [])]
    for gate_id in expected_gate_ids:
        if gate_id and gate_id not in table_text:
            issues.append(f"gate is missing: {gate_id}")
    for token in ("845 total, 844 pass, 0 fail, 1 skip", "124/124 탐지", "111/111", "12/12"):
        if token not in table_text:
            issues.append(f"validation token is missing: {token}")

    footer_text = "\n".join(
        "".join(etree.fromstring(xml).xpath(".//w:t/text()", namespaces=NS))
        for xml in footer_parts
    )
    if "JARVIS 구조 개선 검증" not in footer_text:
        issues.append("footer label is missing")
    if not any(b" PAGE " in xml for xml in footer_parts):
        issues.append("page-number field is missing")

    result = {
        "status": "passed" if not issues else "needs-review",
        "path": str(docx_path),
        "sha256": hashlib.sha256(docx_bytes).hexdigest(),
        "sizeBytes": len(docx_bytes),
        "zipIntegrity": "passed" if not corrupt_member else "failed",
        "pageSizeCm": [round(doc.sections[0].page_width.cm, 2), round(doc.sections[0].page_height.cm, 2)],
        "tableCount": len(doc.tables),
        "appendixCommandCount": len(appendix_ids),
        "gateCount": len(expected_gate_ids),
        "trackedChangeCounts": tracked,
        "issues": issues,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))
    if issues:
        raise SystemExit(1)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Usage: qa-impossible-workflow-30-docx-structure.py <docx> <report.json> <output.json>")
    main(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve(), Path(sys.argv[3]).resolve())
